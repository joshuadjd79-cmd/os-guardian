"""
Compiles OS_Guardian_Complete_Working_Model_Explanation.md to:
1. docs/OS_Guardian_Complete_Working_Model_Explanation.docx
2. docs/OS_Guardian_Complete_Working_Model_Explanation.pdf
"""

import os
import re
import sys
from pathlib import Path

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Image as RLImage, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfgen import canvas

BASE_DIR = Path(r"c:\Users\Steven\OneDrive\Desktop\os guardian")
DOCS_DIR = BASE_DIR / "docs"
SCREENSHOTS_DIR = DOCS_DIR / "screenshots"

MD_PATH = DOCS_DIR / "OS_Guardian_Complete_Working_Model_Explanation.md"
DOCX_PATH = DOCS_DIR / "OS_Guardian_Complete_Working_Model_Explanation.docx"
PDF_PATH = DOCS_DIR / "OS_Guardian_Complete_Working_Model_Explanation.pdf"

# Numbered Canvas for PDF
class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count):
        if self._pageNumber == 1:
            return  # Suppress headers/footers on cover page
        
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748b"))
        
        # Header
        self.drawString(54, 750, "OS Guardian 2.0 — Complete Working Model Explanation & Engineering Defense Guide")
        self.setStrokeColor(colors.HexColor("#cbd5e1"))
        self.setLineWidth(0.5)
        self.line(54, 744, 558, 744)
        
        # Footer
        self.line(54, 45, 558, 45)
        page_str = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(558, 32, page_str)
        self.drawString(54, 32, "Confidential — For Academic Project Review & Examination Defense Only")
        self.restoreState()


def clean_xml_str(val: str) -> str:
    if not val:
        return ""
    # Filter out ASCII control characters below 32 except tab, newline, carriage return
    cleaned = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', val)
    return cleaned


def compile_docx(md_text: str):
    print("Compiling DOCX...")
    doc = Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    lines = [clean_xml_str(l) for l in md_text.splitlines()]
    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    def flush_table():
        nonlocal table_lines
        if not table_lines:
            return
        # Parse table lines
        header_row = [c.strip() for c in table_lines[0].split('|')[1:-1]]
        data_rows = []
        for l in table_lines[2:]:
            if '|' in l:
                data_rows.append([c.strip() for c in l.split('|')[1:-1]])
        
        if header_row and data_rows:
            table = doc.add_table(rows=1 + len(data_rows), cols=len(header_row))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.style = 'Light Shading Accent 1'
            
            # Header
            hdr_cells = table.rows[0].cells
            for idx, name in enumerate(header_row):
                if idx < len(hdr_cells):
                    hdr_cells[idx].text = name
                    for p in hdr_cells[idx].paragraphs:
                        p.runs[0].font.bold = True
                        p.runs[0].font.size = Pt(9.5)
            
            # Data
            for r_idx, row_data in enumerate(data_rows):
                row_cells = table.rows[r_idx + 1].cells
                for c_idx, val in enumerate(row_data):
                    if c_idx < len(row_cells):
                        row_cells[c_idx].text = val
                        for p in row_cells[c_idx].paragraphs:
                            p.runs[0].font.size = Pt(8.5)
            
            doc.add_paragraph()  # Spacing
        table_lines = []

    for line in lines:
        stripped = line.strip()

        # Handle code blocks
        if stripped.startswith('```'):
            if in_code_block:
                in_code_block = False
                p = doc.add_paragraph("\n".join(code_lines))
                p.paragraph_format.left_indent = Inches(0.3)
                for run in p.runs:
                    run.font.name = 'Consolas'
                    run.font.size = Pt(8.5)
                    run.font.color.rgb = RGBColor(0x1e, 0x29, 0x3b)
                code_lines = []
            else:
                in_code_block = True
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        # Handle tables
        if stripped.startswith('|') and stripped.endswith('|'):
            table_lines.append(stripped)
            continue
        elif table_lines:
            flush_table()

        # Headers
        if stripped.startswith('# '):
            h = doc.add_heading(stripped[2:], level=1)
            h.runs[0].font.size = Pt(22)
            h.runs[0].font.color.rgb = RGBColor(0x0f, 0x17, 0x2a)
            h.runs[0].font.bold = True
        elif stripped.startswith('## '):
            h = doc.add_heading(stripped[3:], level=2)
            h.runs[0].font.size = Pt(16)
            h.runs[0].font.color.rgb = RGBColor(0x1e, 0x3a, 0x8a)
            h.runs[0].font.bold = True
        elif stripped.startswith('### '):
            h = doc.add_heading(stripped[4:], level=3)
            h.runs[0].font.size = Pt(13)
            h.runs[0].font.color.rgb = RGBColor(0x33, 0x41, 0x55)
            h.runs[0].font.bold = True
        elif stripped.startswith('#### '):
            h = doc.add_heading(stripped[5:], level=4)
            h.runs[0].font.size = Pt(11)
            h.runs[0].font.color.rgb = RGBColor(0x47, 0x55, 0x69)
        elif stripped.startswith('![') and '](' in stripped and stripped.endswith(')'):
            # Image
            match = re.search(r'\!\[(.*?)\]\((.*?)\)', stripped)
            if match:
                alt = match.group(1)
                img_path = match.group(2)
                full_img_path = DOCS_DIR / img_path
                if full_img_path.exists():
                    p_img = doc.add_paragraph()
                    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_img.paragraph_format.space_before = Pt(8)
                    p_img.paragraph_format.space_after = Pt(2)
                    doc.add_picture(str(full_img_path), width=Inches(6.2))
                    p_cap = doc.add_paragraph(f"Caption: {alt}")
                    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_cap.runs[0].font.italic = True
                    p_cap.runs[0].font.size = Pt(8.5)
                    p_cap.runs[0].font.color.rgb = RGBColor(0x64, 0x74, 0x8b)
        elif stripped.startswith('*Figure ') and stripped.endswith('*'):
            p = doc.add_paragraph(stripped[1:-1])
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(10)
            for run in p.runs:
                run.font.italic = True
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
        elif stripped.startswith('> '):
            p = doc.add_paragraph(stripped[2:])
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.right_indent = Inches(0.4)
            for run in p.runs:
                run.font.italic = True
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        elif stripped.startswith('- ') or stripped.startswith('* '):
            p = doc.add_paragraph(stripped[2:], style='List Bullet')
            for run in p.runs:
                run.font.size = Pt(10)
        elif re.match(r'^\d+\.\s', stripped):
            p = doc.add_paragraph(re.sub(r'^\d+\.\s', '', stripped), style='List Number')
            for run in p.runs:
                run.font.size = Pt(10)
        elif stripped == '---':
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
        elif stripped:
            p = doc.add_paragraph(stripped)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(5)
            for run in p.runs:
                run.font.size = Pt(10)

    if table_lines:
        flush_table()

    doc.save(str(DOCX_PATH))
    print(f"DOCX successfully compiled: {DOCX_PATH} ({os.path.getsize(DOCX_PATH)} bytes)")


def compile_pdf(md_text: str):
    print("Compiling PDF...")
    doc = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=60,
        bottomMargin=54
    )
    
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#0f172a'),
        spaceAfter=8,
        spaceBefore=0
    )
    
    subtitle_style = ParagraphStyle(
        'DocSubTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        textColor=colors.HexColor('#1e3a8a'),
        spaceAfter=14
    )
    
    h1_style = ParagraphStyle(
        'H1',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=15,
        leading=18,
        textColor=colors.HexColor('#0f172a'),
        spaceBefore=16,
        spaceAfter=8,
        keepWithNext=True
    )
    
    h2_style = ParagraphStyle(
        'H2',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=15,
        textColor=colors.HexColor('#1e40af'),
        spaceBefore=12,
        spaceAfter=6,
        keepWithNext=True
    )
    
    h3_style = ParagraphStyle(
        'H3',
        parent=styles['Heading3'],
        fontName='Helvetica-Bold',
        fontSize=10.5,
        leading=13,
        textColor=colors.HexColor('#334155'),
        spaceBefore=10,
        spaceAfter=4,
        keepWithNext=True
    )
    
    body_style = ParagraphStyle(
        'Body',
        parent=styles['BodyText'],
        fontName='Helvetica',
        fontSize=9,
        leading=12.5,
        textColor=colors.HexColor('#1e293b'),
        spaceAfter=5
    )
    
    bullet_style = ParagraphStyle(
        'Bullet',
        parent=body_style,
        leftIndent=15,
        firstLineIndent=-10,
        spaceAfter=3
    )
    
    caption_style = ParagraphStyle(
        'Caption',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=7.5,
        leading=10,
        textColor=colors.HexColor('#475569'),
        alignment=1,  # Center
        spaceBefore=3,
        spaceAfter=8
    )
    
    callout_style = ParagraphStyle(
        'Callout',
        parent=body_style,
        fontName='Helvetica-Oblique',
        fontSize=8.5,
        leading=12,
        textColor=colors.HexColor('#334155'),
        leftIndent=20,
        rightIndent=20,
        spaceBefore=6,
        spaceAfter=6
    )

    code_style = ParagraphStyle(
        'CodeStyle',
        parent=styles['Normal'],
        fontName='Courier',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor('#0f172a'),
        leftIndent=12,
        spaceAfter=2
    )

    story = []
    lines = md_text.splitlines()
    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    def clean_text_for_rl(txt: str) -> str:
        txt = txt.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
        # Bold **text**
        txt = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', txt)
        # Italic *text*
        txt = re.sub(r'\*(.*?)\*', r'<i>\1</i>', txt)
        # Code `text`
        txt = re.sub(r'`(.*?)`', r'<font name="Courier">\1</font>', txt)
        return txt

    def flush_pdf_table():
        nonlocal table_lines
        if not table_lines:
            return
        header_row = [c.strip() for c in table_lines[0].split('|')[1:-1]]
        data_rows = []
        for l in table_lines[2:]:
            if '|' in l:
                data_rows.append([c.strip() for c in l.split('|')[1:-1]])
        
        if header_row and data_rows:
            formatted_data = []
            hdr_cells = [Paragraph(f"<b>{clean_text_for_rl(h)}</b>", caption_style) for h in header_row]
            formatted_data.append(hdr_cells)
            
            for row in data_rows:
                r_cells = []
                for idx, c in enumerate(row):
                    p_cell = Paragraph(clean_text_for_rl(c), ParagraphStyle('TCell', parent=body_style, fontSize=7.5, leading=9.5))
                    r_cells.append(p_cell)
                formatted_data.append(r_cells)
            
            # Auto column widths
            num_cols = len(header_row)
            col_width = 504.0 / num_cols
            t = Table(formatted_data, colWidths=[col_width] * num_cols)
            t.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f1f5f9')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.HexColor('#0f172a')),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
                ('BOX', (0, 0), (-1, -1), 0.8, colors.HexColor('#94a3b8')),
                ('TOPPADDING', (0, 0), (-1, -1), 4),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ]))
            story.append(Spacer(1, 4))
            story.append(t)
            story.append(Spacer(1, 6))
        table_lines = []

    for line in lines:
        stripped = line.strip()

        # Handle code blocks
        if stripped.startswith('```'):
            if in_code_block:
                in_code_block = False
                clean_code = "\n".join(code_lines)
                code_table_data = [[Paragraph(clean_text_for_rl(cl), code_style) for cl in code_lines if cl.strip()]]
                if code_lines:
                    t_box = Table(
                        [[Paragraph("<br/>".join([clean_text_for_rl(cl) for cl in code_lines]), code_style)]],
                        colWidths=[504]
                    )
                    t_box.setStyle(TableStyle([
                        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#f8fafc')),
                        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
                        ('LEFTPADDING', (0, 0), (-1, -1), 8),
                        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
                        ('TOPPADDING', (0, 0), (-1, -1), 6),
                        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
                    ]))
                    story.append(Spacer(1, 4))
                    story.append(t_box)
                    story.append(Spacer(1, 6))
                code_lines = []
            else:
                in_code_block = True
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        # Tables
        if stripped.startswith('|') and stripped.endswith('|'):
            table_lines.append(stripped)
            continue
        elif table_lines:
            flush_pdf_table()

        if stripped.startswith('# '):
            story.append(Paragraph(clean_text_for_rl(stripped[2:]), title_style))
        elif stripped.startswith('## '):
            if "EXECUTIVE TABLE OF CONTENTS" in stripped or "Chapter 1:" in stripped:
                story.append(PageBreak())
            story.append(Paragraph(clean_text_for_rl(stripped[3:]), h1_style))
        elif stripped.startswith('### '):
            story.append(Paragraph(clean_text_for_rl(stripped[4:]), h2_style))
        elif stripped.startswith('#### '):
            story.append(Paragraph(clean_text_for_rl(stripped[5:]), h3_style))
        elif stripped.startswith('![') and '](' in stripped and stripped.endswith(')'):
            match = re.search(r'\!\[(.*?)\]\((.*?)\)', stripped)
            if match:
                alt = match.group(1)
                img_path = match.group(2)
                full_img_path = DOCS_DIR / img_path
                if full_img_path.exists():
                    story.append(Spacer(1, 4))
                    img = RLImage(str(full_img_path), width=500, height=312.5)
                    story.append(img)
        elif stripped.startswith('*Figure ') and stripped.endswith('*'):
            story.append(Paragraph(clean_text_for_rl(stripped[1:-1]), caption_style))
            story.append(Spacer(1, 4))
        elif stripped.startswith('> '):
            story.append(Paragraph(clean_text_for_rl(stripped[2:]), callout_style))
        elif stripped.startswith('- ') or stripped.startswith('* '):
            bullet_txt = "&bull; " + clean_text_for_rl(stripped[2:])
            story.append(Paragraph(bullet_txt, bullet_style))
        elif re.match(r'^\d+\.\s', stripped):
            num_txt = clean_text_for_rl(stripped)
            story.append(Paragraph(num_txt, bullet_style))
        elif stripped == '---':
            story.append(Spacer(1, 4))
            story.append(HRFlowable(width="100%", thickness=0.8, color=colors.HexColor('#cbd5e1'), spaceBefore=4, spaceAfter=6))
        elif stripped:
            story.append(Paragraph(clean_text_for_rl(stripped), body_style))

    if table_lines:
        flush_pdf_table()

    doc.build(story, canvasmaker=NumberedCanvas)
    print(f"PDF successfully compiled: {PDF_PATH} ({os.path.getsize(PDF_PATH)} bytes)")


if __name__ == "__main__":
    with open(MD_PATH, "r", encoding="utf-8") as f:
        md_text = f.read()
    compile_docx(md_text)
    compile_pdf(md_text)
